#!/usr/bin/env python3
"""Exporta reglas desde control-engorda-gac.xlsx a src/config/*.json"""
import json
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "docs/plantillas/control-engorda-gac.xlsx"
GAM = ROOT / "docs/plantillas/gam-tesoreria-2025.xlsx"
OUT = ROOT / "src/config"
CANTIDAD_BASE = 18500

def json_default(obj):
    if isinstance(obj, (datetime, date)):
        return obj.isoformat()
    if isinstance(obj, Decimal):
        return float(obj)
    return str(obj)

def export_parametros(wb):
    ws = wb["Parametros"]
    rows = []
    for r in range(3, ws.max_row + 1):
        desc = ws.cell(r, 2).value
        if not desc:
            continue
        rows.append({
            "descripcion": str(desc).strip(),
            "marca": ws.cell(r, 3).value,
            "proveedor": ws.cell(r, 4).value,
            "presentacion_kg": ws.cell(r, 5).value,
            "precio_bulto": float(ws.cell(r, 6).value or 0),
            "precio_kg": float(ws.cell(r, 7).value or 0),
            "nota": ws.cell(r, 8).value,
        })
    return rows

def export_tabla_diaria(wb):
    ws = wb["Sheet2"]
    filas = []
    for r in range(3, ws.max_row + 1):
        dia = ws.cell(r, 2).value
        if dia is None or str(dia).strip().lower() == "total":
            break
        try:
            dia_int = int(dia)
        except (TypeError, ValueError):
            continue
        peso_g = ws.cell(r, 16).value
        tipo = ws.cell(r, 17).value
        biomasa = ws.cell(r, 18).value
        tasa = ws.cell(r, 19).value
        if peso_g is None and tipo is None:
            continue
        biomasa_f = float(biomasa or 0)
        tasa_f = float(tasa or 0)
        filas.append({
            "dia": dia_int,
            "peso_promedio_g": round(float(peso_g or 0), 6),
            "tipo_alimento": str(tipo or "").strip(),
            "biomasa_base_kg": round(biomasa_f, 6),
            "tasa_alimentacion_pct": round(tasa_f, 6),
            "kg_alimento_base_dia": round(biomasa_f * tasa_f, 6),
        })
    return {"cantidad_base": CANTIDAD_BASE, "filas": filas}

def export_kpi_labels(wb):
    ws = wb["1-2025"]
    labels = []
    for r in range(4, 30):
        label = ws.cell(r, 12).value
        if not label:
            continue
        s = str(label).strip()
        if s == "No. de Ciclo":
            continue
        labels.append(s)
        if s.startswith("Dias Transcurridos del Ciclo"):
            break
    return labels

def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "parametros-alimento.json").write_text(
        json.dumps(export_parametros(wb), ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )
    (OUT / "tabla-alimentacion-diaria.json").write_text(
        json.dumps(export_tabla_diaria(wb), ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )
    (OUT / "ciclo-engorda-kpis.json").write_text(
        json.dumps({"labels": export_kpi_labels(wb)}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    wb.close()
    print("Plantillas exportadas a", OUT)

if __name__ == "__main__":
    main()
